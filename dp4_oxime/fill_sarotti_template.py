#!/usr/bin/env python3
"""
Fill user-provided Sarotti DP4+ template (decrypted .xlsx).

Usage:
  python3 fill_sarotti_template.py /path/to/dp4_plus.xlsx

Output: dp4_plus_DOLDURULMUS.xlsx next to input (or in dp4_oxime/)
"""

from __future__ import annotations

import sys
from copy import copy
from pathlib import Path

import openpyxl

from build_dp4_excel import EXP_C, EXP_H, _avg
from parse_jaguar_nmr import parse_jaguar_shieldings

DIR = Path(__file__).parent


def zone_b_rows(shield_e, shield_z):
    rows = []
    for exp, _role, sp2, le, lz in EXP_H:
        rows.append(
            ("H", "x" if sp2 else None, exp, _avg(shield_e, le), _avg(shield_z, lz))
        )
    for exp, _role, sp2, le, lz in EXP_C:
        rows.append(("C", "x" if sp2 else None, exp, shield_e[le], shield_z[lz]))
    return rows


def find_zone_b_start(ws) -> tuple[int, dict[str, int]] | None:
    """Locate header row with Nuclei / Experimental / Isomer."""
    for r in range(1, min(ws.max_row, 80) + 1):
        vals = {ws.cell(r, c).value: c for c in range(1, 15) if ws.cell(r, c).value}
        keys = " ".join(str(k).lower() for k in vals if k)
        if "nuclei" in keys and "experimental" in keys:
            cols = {}
            for k, c in vals.items():
                ks = str(k).lower()
                if "nuclei" in ks:
                    cols["nuclei"] = c
                elif "sp2" in ks:
                    cols["sp2"] = c
                elif "experimental" in ks:
                    cols["exp"] = c
                elif "isomer 1" in ks or ks.strip() == "1":
                    cols["iso1"] = c
                elif "isomer 2" in ks or ks.strip() == "2":
                    cols["iso2"] = c
            if "nuclei" in cols and "exp" in cols and "iso1" in cols:
                return r, cols
    return None


def main() -> None:
    if len(sys.argv) < 2:
        print("Usage: python3 fill_sarotti_template.py YOUR_dp4_plus.xlsx")
        sys.exit(1)

    src = Path(sys.argv[1])
    if not src.exists():
        print(f"File not found: {src}")
        sys.exit(1)

    shield_e = parse_jaguar_shieldings(DIR / "jag_G5_E_spe.log")
    shield_z = parse_jaguar_shieldings(DIR / "jag_G1_spe.log")
    rows = zone_b_rows(shield_e, shield_z)

    wb = openpyxl.load_workbook(src)
    if "Main" not in wb.sheetnames:
        print("Sheet 'Main' not found; sheets:", wb.sheetnames)
        sys.exit(1)
    ws = wb["Main"]
    found = find_zone_b_start(ws)
    if not found:
        print("Could not find Zone B header row. Use DP4_PLUS_HAKEME_DOLDURULMUS.xlsx instead.")
        sys.exit(1)

    hr, cols = found
    data_start = hr + 1
    for i, (nuc, sp2, exp, s1, s2) in enumerate(rows):
        r = data_start + i
        ws.cell(r, cols["nuclei"], nuc)
        if "sp2" in cols:
            ws.cell(r, cols["sp2"], sp2 if sp2 else "")
        ws.cell(r, cols["exp"], exp)
        ws.cell(r, cols["iso1"], round(s1, 4))
        ws.cell(r, cols["iso2"], round(s2, 4))

    out = src.parent / f"{src.stem}_DOLDURULMUS{src.suffix}"
    wb.save(out)
    print(f"Filled {len(rows)} rows → {out}")
    print("Open in Excel; check Zone A (B3LYP, PCM, 6-31G**, Shielding Tensors).")


if __name__ == "__main__":
    main()
